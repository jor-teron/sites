#!/usr/bin/env python3
"""Builds merge_sample.xlsx (Draftly MM3): 2 sheets, date / number-format / boolean columns
and formulas. openpyxl writes formulas without results, so the file is passed through
LibreOffice (headless) which stores the calculated values. Needs: python3 openpyxl, soffice.
Usage: python3 make_merge_sample.py   (writes merge_sample.xlsx next to this script)"""
import datetime, os, shutil, subprocess, tempfile
import openpyxl

here = os.path.dirname(os.path.abspath(__file__))
wb = openpyxl.Workbook()
c = wb.active
c.title = 'Customers'
c.append(['Name', 'No', 'Company', 'City', 'Joined', 'Amount', 'Discount', 'Total', 'Active', 'Note'])
rows = [
    ['Anna Smith', 1001, 'Smith & Sons, Ltd.', 'Paris', datetime.date(2024, 3, 5), 1250.0, 0.10, True, 'Prefers email.\nCall after 5 pm.'],
    ['Rahul Mehta', 1002, 'Mehta Traders', 'Mumbai', datetime.date(2023, 11, 20), 980.5, 0.0, True, None],
    ['Zoë Müller', 1003, 'Müller "Werkzeug" GmbH', 'Köln', datetime.date(2025, 1, 15), 2400.0, 0.05, False, 'New customer'],
]
for i, r in enumerate(rows, start=2):
    c.append(r[:7] + [None] + r[7:])
    c.cell(i, 8).value = '=F%d*(1-G%d)' % (i, i)          # Total (formula)
    c.cell(i, 5).number_format = 'dd-mmm-yyyy'
    c.cell(i, 6).number_format = '#,##0.00'
    c.cell(i, 7).number_format = '0%'
    c.cell(i, 8).number_format = '#,##0.00'
o = wb.create_sheet('Orders')
o.append(['Order', 'No', 'Item', 'Qty', 'Unit price', 'Line total', 'Ordered'])
orders = [['A-100', 1001, 'Widget', 12, 4.25, datetime.date(2026, 9, 30)],
          ['A-101', 1002, 'Gadget "Pro"', 3, 129.9, datetime.date(2026, 10, 2)],
          ['A-102', 1003, 'Spanner', 40, 7.5, datetime.date(2026, 10, 12)]]
for i, r in enumerate(orders, start=2):
    o.append(r[:5] + [None, r[5]])
    o.cell(i, 6).value = '=D%d*E%d' % (i, i)
    o.cell(i, 5).number_format = '#,##0.00'
    o.cell(i, 6).number_format = '#,##0.00'
    o.cell(i, 7).number_format = 'yyyy-mm-dd'
for ws, widths in ((c, [16, 7, 26, 10, 13, 11, 10, 11, 8, 26]), (o, [9, 7, 16, 6, 11, 11, 12])):
    for k, w in enumerate(widths):
        ws.column_dimensions[chr(65 + k)].width = w
tmp = tempfile.mkdtemp()
try:
    src = os.path.join(tmp, 'in', 'merge_sample.xlsx')
    os.makedirs(os.path.dirname(src))
    wb.save(src)
    subprocess.check_call(['soffice', '--headless', '--convert-to', 'xlsx', '--outdir', os.path.join(tmp, 'out'), src],
                          stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    shutil.copy(os.path.join(tmp, 'out', 'merge_sample.xlsx'), os.path.join(here, 'merge_sample.xlsx'))
finally:
    shutil.rmtree(tmp, ignore_errors=True)
print('wrote', os.path.join(here, 'merge_sample.xlsx'))
